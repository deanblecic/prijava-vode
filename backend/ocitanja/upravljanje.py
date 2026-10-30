"""Web sučelje za administratore: mjesečni pregled, provjera očitanja, stanovi i izvoz.

Svaki administrator radi u jednoj, odabranoj zgradi (request.zgrada) i dohvaća samo njezine
podatke. Superadministrator i administrator više zgrada biraju zgradu u zaglavlju.
"""
import calendar
import datetime
from pathlib import Path
from decimal import Decimal
from functools import wraps

from django.conf import settings
from django.contrib import messages
from django.contrib.auth import views as auth_views
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.db.models import Count, Q
from django.http import FileResponse, Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font

from .forms import IspravakForm, PostavkeForm, PromjenaVlasnikaForm, RucniUnosForm
from .models import Korisnik, Ocitanje, Postavke, Stan
from .servisi import (
    GreskaOcitanja,
    ispravi_ocitanje,
    postavi_status,
    predaj_ocitanje,
    promijeni_vlasnika,
    verzija_aplikacije,
)


KLJUC_ZGRADE = "zgrada_id"


def trenutna_zgrada(request):
    """Zgrada odabrana u ovoj sesiji, ako je korisnik smije vidjeti.

    Administrator samo jedne zgrade uvijek radi u njoj; superadministrator mora odabrati.
    """
    if not request.user.is_authenticated or not request.user.je_administrator:
        return None
    zgrade = request.user.zgrade_za_upravljanje()
    odabrana = request.session.get(KLJUC_ZGRADE)
    if odabrana:
        zgrada = zgrade.filter(pk=odabrana).first()
        if zgrada:
            return zgrada
        request.session.pop(KLJUC_ZGRADE, None)
    if not request.user.is_superuser and len(zgrade) == 1:
        return zgrade[0]
    return None


def samo_administrator(view):
    @login_required
    @wraps(view)
    def omot(request, *args, **kwargs):
        if not request.user.je_administrator:
            raise PermissionDenied
        return view(request, *args, **kwargs)

    return omot


def za_zgradu(view):
    """Administratorska stranica koja radi s jednom zgradom; bez odabrane vodi na popis zgrada."""

    @samo_administrator
    @wraps(view)
    def omot(request, *args, **kwargs):
        request.zgrada = trenutna_zgrada(request)
        if request.zgrada is None:
            return redirect("zgrade")
        return view(request, *args, **kwargs)

    return omot


def kontekst(request):
    """Odabrana zgrada i izbornik zgrada u zaglavlju svake stranice."""
    if not request.user.is_authenticated or not request.user.je_administrator:
        return {}
    zgrade = list(request.user.zgrade_za_upravljanje())
    return {
        "zgrada": getattr(request, "zgrada", None) or trenutna_zgrada(request),
        "moje_zgrade": zgrade if request.user.is_superuser or len(zgrade) > 1 else [],
    }


class Prijava(auth_views.LoginView):
    template_name = "upravljanje/prijava.html"
    redirect_authenticated_user = True

    def form_valid(self, form):
        # Web je za sada samo za administratore; suvlasnici očitanje unose u aplikaciji.
        if not form.get_user().je_administrator:
            return self.render_to_response(self.get_context_data(form=form, samo_aplikacija=True))
        return super().form_valid(form)


@samo_administrator
def zgrade(request):
    mjesec = timezone.localdate().replace(day=1)
    popis = []
    for zgrada in request.user.zgrade_za_upravljanje().annotate(broj_stanova=Count("stanovi")):
        redovi = redovi_mjeseca(zgrada, mjesec)
        popis.append({
            "zgrada": zgrada,
            "broj_stanova": zgrada.broj_stanova,
            "za_ocitati": len(redovi),
            "neocitano": sum(1 for r in redovi if not r["ocitanje"] or r["ocitanje"].status == Ocitanje.Status.ODBIJENO),
        })
    return render(request, "upravljanje/zgrade.html", {"popis": popis, "mjesec": mjesec})


@samo_administrator
def odaberi_zgradu(request):
    if request.method != "POST":
        return redirect("zgrade")
    zgrada = request.user.zgrade_za_upravljanje().filter(pk=request.POST.get("zgrada") or 0).first()
    if zgrada is None:
        request.session.pop(KLJUC_ZGRADE, None)
        return redirect("zgrade")
    request.session[KLJUC_ZGRADE] = zgrada.pk
    return redirect("pregled")


def mjesec_iz_upita(request):
    try:
        return datetime.datetime.strptime(request.GET.get("mjesec", ""), "%Y-%m").date()
    except ValueError:
        return timezone.localdate().replace(day=1)


def pomakni_mjesec(mjesec, za):
    indeks = mjesec.year * 12 + mjesec.month - 1 + za
    return datetime.date(indeks // 12, indeks % 12 + 1, 1)


def zadnji_dan(mjesec):
    return mjesec.replace(day=calendar.monthrange(mjesec.year, mjesec.month)[1])


def vlasnici_u_mjesecu(stan, mjesec):
    veze = stan.vlasnistva.filter(od_datuma__lte=zadnji_dan(mjesec)).filter(
        Q(do_datuma__isnull=True) | Q(do_datuma__gte=mjesec)
    ).select_related("korisnik")
    return [v.korisnik for v in veze]


def redovi_mjeseca(zgrada, mjesec):
    ocitanja = {
        o.vodomjer.stan_id: o
        for o in Ocitanje.objects.filter(razdoblje=mjesec, vodomjer__stan__zgrada=zgrada)
        .select_related("vodomjer__stan", "korisnik")
    }
    redovi = []
    # Stanovi bez aktivnog vodomjera (npr. Grad Rijeka) ne očitavaju se, pa nisu u pregledu.
    stanovi = zgrada.stanovi.filter(vodomjeri__aktivan=True).distinct().prefetch_related("vlasnistva__korisnik")
    for stan in stanovi:
        redovi.append({
            "stan": stan,
            "vlasnici": ", ".join(k.get_full_name() or k.username for k in vlasnici_u_mjesecu(stan, mjesec)),
            "ocitanje": ocitanja.get(stan.id),
        })
    return sorted(redovi, key=lambda r: kljuc_stana(r["stan"].broj_stana))


def kljuc_stana(broj):
    """Prirodni redoslijed: 2 prije 10, 2A iza 2."""
    znamenke = "".join(c for c in broj if c.isdigit())
    return (int(znamenke) if znamenke else 0, broj)


@za_zgradu
def pregled(request):
    mjesec = mjesec_iz_upita(request)
    redovi = redovi_mjeseca(request.zgrada, mjesec)
    predana = [r["ocitanje"] for r in redovi if r["ocitanje"]]
    vazeca = [o for o in predana if o.status != Ocitanje.Status.ODBIJENO]
    return render(request, "upravljanje/pregled.html", {
        "mjesec": mjesec,
        "prethodni": pomakni_mjesec(mjesec, -1),
        "sljedeci": pomakni_mjesec(mjesec, 1),
        "redovi": redovi,
        "broj_predanih": len(predana),
        "broj_stanova": len(redovi),
        "za_provjeru": sum(1 for o in predana if o.za_provjeru),
        "nepotvrdena": sum(1 for o in predana if o.status == Ocitanje.Status.PREDANO),
        "ukupna_razlika": sum((o.razlika for o in vazeca), Decimal(0)),
    })


@za_zgradu
def ocitanje_detalj(request, pk):
    ocitanje = get_object_or_404(
        Ocitanje.objects.select_related("vodomjer__stan", "korisnik"), pk=pk, vodomjer__stan__zgrada=request.zgrada
    )
    forma = IspravakForm(initial={"novo_stanje": ocitanje.novo_stanje, "napomena": ocitanje.napomena})
    if request.method == "POST":
        akcija = request.POST.get("akcija")
        upozorenje = None
        if akcija == "potvrdi":
            upozorenje = postavi_status(ocitanje, Ocitanje.Status.POTVRDENO)
            messages.success(request, f"Očitanje za stan {ocitanje.vodomjer.stan.broj_stana} je potvrđeno.")
        elif akcija == "odbij":
            upozorenje = postavi_status(ocitanje, Ocitanje.Status.ODBIJENO)
            messages.info(request, "Očitanje je odbijeno; suvlasnik ga može ponovno predati.")
        elif akcija == "ispravi":
            forma = IspravakForm(request.POST)
            if forma.is_valid():
                upozorenje = ispravi_ocitanje(
                    ocitanje, forma.cleaned_data["novo_stanje"], forma.cleaned_data["napomena"]
                )
                messages.success(request, "Očitanje je ispravljeno.")
            else:
                return render(request, "upravljanje/ocitanje.html", {"o": ocitanje, "forma": forma})
        if upozorenje:
            messages.warning(request, upozorenje)
        return redirect(f"{reverse('pregled')}?mjesec={ocitanje.razdoblje:%Y-%m}")
    return render(request, "upravljanje/ocitanje.html", {
        "o": ocitanje,
        "forma": forma,
        "vlasnici": vlasnici_u_mjesecu(ocitanje.vodomjer.stan, ocitanje.razdoblje),
    })


@za_zgradu
def stan_detalj(request, pk):
    stan = get_object_or_404(Stan, pk=pk, zgrada=request.zgrada)
    vodomjer = stan.vodomjeri.filter(aktivan=True).first()
    mjesec = mjesec_iz_upita(request)
    unos = RucniUnosForm(initial={"datum_ocitanja": min(zadnji_dan(mjesec), timezone.localdate())})
    if request.method == "POST":
        unos = RucniUnosForm(request.POST, request.FILES)
        if vodomjer and unos.is_valid():
            d = unos.cleaned_data
            trenutak = timezone.make_aware(datetime.datetime.combine(d["datum_ocitanja"], datetime.time(12)))
            try:
                predaj_ocitanje(
                    request.user, vodomjer, d["novo_stanje"], datum_ocitanja=trenutak,
                    fotografija=d["fotografija"], napomena=d["napomena"] or "Unio administrator",
                    potvrdi_manje=request.POST.get("potvrdi_manje") == "1",
                )
            except GreskaOcitanja as greska:
                messages.error(request, greska.poruka)
            else:
                messages.success(request, "Očitanje je spremljeno.")
                return redirect(f"{reverse('pregled')}?mjesec={trenutak:%Y-%m}")
    return render(request, "upravljanje/stan.html", {
        "stan": stan,
        "vodomjer": vodomjer,
        "vlasnistva": stan.vlasnistva.select_related("korisnik").order_by("-od_datuma"),
        "ocitanja": Ocitanje.objects.filter(vodomjer__stan=stan).select_related("korisnik", "vodomjer"),
        "unos": unos,
    })


@za_zgradu
def promjena_vlasnika(request, pk):
    stan = get_object_or_404(Stan, pk=pk, zgrada=request.zgrada)
    forma = PromjenaVlasnikaForm(
        request.POST or None, zgrada=request.zgrada, initial={"od_datuma": timezone.localdate()}
    )
    if request.method == "POST" and forma.is_valid():
        d = forma.cleaned_data
        stari = vlasnici_u_mjesecu(stan, timezone.localdate().replace(day=1))
        novi = list(d["postojeci"])
        if d["korisnicko_ime"]:
            novi.append(Korisnik.objects.create_user(
                d["korisnicko_ime"], password=d["lozinka"], first_name=d["ime"], last_name=d["prezime"],
            ))
        promijeni_vlasnika(stan, novi, d["od_datuma"])
        if d["deaktiviraj_stare"]:
            for korisnik in stari:
                if korisnik not in novi and not korisnik.je_administrator and not korisnik.vlasnistva.filter(
                    Q(do_datuma__isnull=True) | Q(do_datuma__gte=d["od_datuma"])
                ).exists():
                    korisnik.is_active = False
                    korisnik.save(update_fields=["is_active"])
        imena = ", ".join(k.get_full_name() or k.username for k in novi)
        messages.success(request, f"Stan {stan.broj_stana} od {d['od_datuma']:%d.%m.%Y.} pripada: {imena}.")
        return redirect("stan", pk=stan.pk)
    return render(request, "upravljanje/promjena_vlasnika.html", {"stan": stan, "forma": forma})


@za_zgradu
def izvoz_excel(request):
    mjesec = mjesec_iz_upita(request)
    knjiga = Workbook()
    list_ = knjiga.active
    list_.title = f"{mjesec:%m-%Y}"
    zaglavlje = [
        "Broj stana", "Ime i prezime", "Datum očitanja",
        "Prethodno stanje (m³)", "Novo stanje (m³)", "Razlika (m³)", "Status",
    ]
    list_.append(zaglavlje)
    for celija in list_[1]:
        celija.font = Font(bold=True)
    for red in redovi_mjeseca(request.zgrada, mjesec):
        o = red["ocitanje"]
        if o:
            list_.append([
                red["stan"].broj_stana, red["vlasnici"], timezone.localtime(o.datum_ocitanja).date(),
                o.prethodno_stanje, o.novo_stanje, o.razlika, o.get_status_display(),
            ])
        else:
            list_.append([red["stan"].broj_stana, red["vlasnici"], None, None, None, None, "Nije predano"])
    zadnji = list_.max_row
    list_.append(["Ukupno", None, None, None, None, f"=SUMIF(G2:G{zadnji},\"<>Odbijeno\",F2:F{zadnji})", None])
    list_.cell(row=zadnji + 1, column=1).font = Font(bold=True)
    for red in list_.iter_rows(min_row=2, min_col=3, max_col=3):
        red[0].number_format = "DD.MM.YYYY"
    for stupac, sirina in zip("ABCDEFG", [11, 28, 15, 20, 17, 13, 14]):
        list_.column_dimensions[stupac].width = sirina
    odgovor = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    odgovor["Content-Disposition"] = f'attachment; filename="ocitanja-{request.zgrada.oznaka}-{mjesec:%Y-%m}.xlsx"'
    knjiga.save(odgovor)
    return odgovor


@samo_administrator
def fotografija(request, putanja):
    """Fotografije vodomjera vidi samo administrator njihove zgrade, pa ih ne poslužuje izravno web server."""
    korijen = Path(settings.MEDIA_ROOT).resolve()
    datoteka = (korijen / putanja).resolve()
    if not datoteka.is_relative_to(korijen) or not datoteka.is_file():
        raise Http404
    if not Ocitanje.objects.filter(
        fotografija=putanja, vodomjer__stan__zgrada__in=request.user.zgrade_za_upravljanje()
    ).exists():
        raise Http404
    return FileResponse(open(datoteka, "rb"))


def aplikacija(request):
    """Javna stranica s koje suvlasnici preuzimaju Android aplikaciju."""
    apk = settings.APK_DATOTEKA
    return render(request, "upravljanje/aplikacija.html", {
        "dostupna": apk.is_file(),
        "velicina_mb": round(apk.stat().st_size / 1_000_000, 1) if apk.is_file() else None,
        "verzija": verzija_aplikacije(),
    })


def aplikacija_apk(request):
    if not settings.APK_DATOTEKA.is_file():
        raise Http404
    return FileResponse(
        open(settings.APK_DATOTEKA, "rb"),
        as_attachment=True,
        filename="ocitanje-vode.apk",
        content_type="application/vnd.android.package-archive",
    )


@za_zgradu
def postavke(request):
    forma = PostavkeForm(request.POST or None, instance=Postavke.ucitaj(request.zgrada))
    if request.method == "POST" and forma.is_valid():
        forma.save()
        messages.success(request, "Postavke su spremljene. Mobiteli ih preuzimaju pri sljedećem otvaranju aplikacije.")
        return redirect("postavke")
    return render(request, "upravljanje/postavke.html", {"forma": forma})
