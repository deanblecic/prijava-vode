import tempfile
from decimal import Decimal
from pathlib import Path

from django.core.management import CommandError, call_command
from django.test import TestCase
from openpyxl import Workbook, load_workbook

from .management.commands.uvoz_stanova import korisnicko_ime, razdvoji_ime
from .models import Korisnik, Stan


class UvozStanovaTest(TestCase):
    def setUp(self):
        self.mapa = Path(tempfile.mkdtemp())
        self.tablica = self.mapa / "voda.xlsx"
        self.izlaz = self.mapa / "pristup.xlsx"
        knjiga = Workbook()
        for red in [
            ("01", "Dzhurenko", 314),
            ("02", "Grad Rijeka", None),
            ("05", "Grad Rijeka", None),
            ("06", "Prelić Gerdević Nataša", 2554),
            ("17", "Njirić Sanja", 1039),
            ("19", "Njirić  Sanja", 1574),
        ]:
            knjiga.active.append(red)
        knjiga.save(self.tablica)

    def uvezi(self):
        call_command("uvoz_stanova", str(self.tablica), pristupni_podaci=str(self.izlaz))

    def test_imena(self):
        self.assertEqual(razdvoji_ime("Prelić Gerdević Nataša"), ("Nataša", "Prelić Gerdević"))
        self.assertEqual(razdvoji_ime("Dzhurenko"), ("", "Dzhurenko"))
        self.assertEqual(korisnicko_ime("Šverko Sara"), "sverko.sara")
        self.assertEqual(korisnicko_ime("Đurić Ivan"), "djuric.ivan")

    def test_uvoz(self):
        self.uvezi()
        self.assertEqual(Stan.objects.count(), 6)
        stan6 = Stan.objects.get(broj_stana="06")
        self.assertEqual(stan6.vodomjeri.get().pocetno_stanje, Decimal("2554"))
        natasa = Korisnik.objects.get(username="prelic.gerdevic.natasa")
        self.assertEqual((natasa.first_name, natasa.last_name), ("Nataša", "Prelić Gerdević"))

        sanja = Korisnik.objects.get(username="njiric.sanja")
        self.assertEqual(sorted(s.broj_stana for s in sanja.aktivni_stanovi()), ["17", "19"])

        grad = Korisnik.objects.get(username="grad.rijeka")
        self.assertEqual(grad.get_full_name(), "Grad Rijeka")
        self.assertFalse(grad.is_active)
        self.assertFalse(grad.has_usable_password())
        self.assertFalse(Stan.objects.get(broj_stana="02").vodomjeri.exists())

        redovi = list(load_workbook(self.izlaz).active.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(redovi), 4)
        sanja_red = next(r for r in redovi if r[2] == "njiric.sanja")
        self.assertEqual(sanja_red[0], "17, 19")
        self.assertTrue(self.client.login(username="njiric.sanja", password=sanja_red[3]))

    def test_ponovni_uvoz_odbijen(self):
        self.uvezi()
        with self.assertRaises(CommandError):
            self.uvezi()
        self.assertEqual(Stan.objects.count(), 6)
